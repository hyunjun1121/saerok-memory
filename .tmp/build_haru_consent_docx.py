from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\project\saerok-memory")
OUTPUT = ROOT / "docs" / "Haru_개인정보_처리_동의서_데모.docx"
STAMP = Path(r"C:\Users\mnb92\Downloads\도장.png")

# standard_business_brief preset + named Korean legibility override.
FONT = "Malgun Gothic"
NAVY = "213B52"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "222222"
MUTED = "64717D"
LIGHT_GRAY = "F2F4F7"
PALE_BLUE = "EEF5FA"
PALE_GOLD = "FFF7E6"
GOLD = "B37A00"
RED = "A61B2B"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(
    run,
    *,
    size: float | None = None,
    bold: bool | None = None,
    color: str | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = FONT
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{key}"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top=65, start=120, bottom=65, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = "D5DCE3", size: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def set_table_geometry(table, widths_dxa: list[int], *, indent_dxa: int = TABLE_INDENT_DXA) -> None:
    assert sum(widths_dxa) == TABLE_WIDTH_DXA
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8.5, color=MUTED)
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def add_heading(doc: Document, text: str, level: int = 1, *, page_break_before: bool = False) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = page_break_before
    run = paragraph.add_run(text)
    set_run_font(run)


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None, after: float = 6) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.18
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, size=11.2, bold=True, color=INK)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest, size=11.2, color=INK)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=11.2, color=INK)


def add_check_line(doc: Document, left: str = "동의합니다", right: str = "동의하지 않습니다") -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = False
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    text = f"□  {left}" if not right else f"□  {left}        □  {right}"
    run = paragraph.add_run(text)
    set_run_font(run, size=12, bold=True, color=NAVY)


def add_callout(doc: Document, label: str, text: str, *, fill: str, accent: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    # Named callout override: table indent follows its 180-DXA start cell margin.
    set_table_geometry(table, [TABLE_WIDTH_DXA], indent_dxa=180)
    set_table_borders(table, color=fill, size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=110, bottom=110, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r1 = p.add_run(label + "  ")
    set_run_font(r1, size=10.5, bold=True, color=accent)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=INK)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_info_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_geometry(table, [1900, 7460])
    set_table_borders(table)
    for row_index, (label, detail) in enumerate(rows):
        label_cell, detail_cell = table.rows[row_index].cells
        label_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        detail_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(label_cell, LIGHT_GRAY)
        p1 = label_cell.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.line_spacing = 1.1
        r1 = p1.add_run(label)
        set_run_font(r1, size=10.5, bold=True, color=NAVY)
        p2 = detail_cell.paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.16
        r2 = p2.add_run(detail)
        set_run_font(r2, size=10.5, color=INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)


def add_easy_promise(doc: Document, title: str, detail: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.12
    title_run = paragraph.add_run(title + "  ")
    set_run_font(title_run, size=11.2, bold=True, color=BLUE)
    detail_run = paragraph.add_run(detail)
    set_run_font(detail_run, size=11.2, color=INK)


def style_document(doc: Document) -> None:
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.different_first_page_header_footer = False
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.right_margin = Inches(1.0)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(1.0)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11.2)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    heading_tokens = {
        "Heading 1": (16, BLUE, 12, 6),
        "Heading 2": (13, BLUE, 10, 5),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    left = hp.add_run("HARU  |  개인정보 처리 안내 및 동의서")
    set_run_font(left, size=8.5, bold=True, color=MUTED)
    hp.add_run("\t")
    right = hp.add_run("DEMO DRAFT")
    set_run_font(right, size=8.5, bold=True, color=RED)
    tabs = hp.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    left_footer = fp.add_run("오가사 LABs · 대표 김현준")
    set_run_font(left_footer, size=8.5, color=MUTED)
    fp.add_run("\t")
    tabs = fp.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    add_page_field(fp)


def build_document() -> Document:
    if not STAMP.exists():
        raise FileNotFoundError(STAMP)

    doc = Document()
    style_document(doc)
    props = doc.core_properties
    props.title = "Haru 개인정보 처리 안내 및 동의서 - 데모"
    props.subject = "고령 이용자의 개인 일기 및 기억 기록 서비스 개인정보 동의서 데모"
    props.author = "오가사 LABs"
    props.keywords = "Haru, 개인정보, 동의서, 데모"

    # Page 1: memo_masthead opening, without decorative bottom rule.
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(8)
    kicker.paragraph_format.space_after = Pt(4)
    kr = kicker.add_run("HARU PRIVACY NOTICE")
    set_run_font(kr, size=10, bold=True, color=GOLD)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    tr = title.add_run("개인정보 처리 안내 및 동의서")
    set_run_font(tr, size=23, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(8)
    sr = subtitle.add_run("개인의 하루와 기억을 편안하게 기록하기 위한 안내")
    set_run_font(sr, size=13, color=MUTED)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(8)
    mr1 = meta.add_run("서비스 운영자  ")
    set_run_font(mr1, size=10.5, bold=True, color=NAVY)
    mr2 = meta.add_run("오가사 LABs    ")
    set_run_font(mr2, size=10.5, color=INK)
    mr3 = meta.add_run("대표자  ")
    set_run_font(mr3, size=10.5, bold=True, color=NAVY)
    mr4 = meta.add_run("김현준    ")
    set_run_font(mr4, size=10.5, color=INK)
    mr5 = meta.add_run("소재 국가  ")
    set_run_font(mr5, size=10.5, bold=True, color=NAVY)
    mr6 = meta.add_run("대한민국")
    set_run_font(mr6, size=10.5, color=INK)

    add_callout(
        doc,
        "데모용 초안",
        "현재 문서는 시연과 화면 구성 확인을 위한 예시입니다. 실제 운영 전 보유기간, 수탁업체, 문의처를 확정해야 합니다.",
        fill=PALE_GOLD,
        accent=GOLD,
    )

    add_heading(doc, "먼저, 쉽게 알려드려요", 1)
    add_easy_promise(doc, "기록과 삭제", "하루의 이야기와 기억을 저장하며, 원하면 언제든 삭제할 수 있습니다.")
    add_easy_promise(doc, "음성", "녹음과 글 변환은 이용자가 원할 때만 사용합니다.")
    add_easy_promise(doc, "공유", "가족이나 보호자에게 자동으로 공유하지 않습니다.")
    add_easy_promise(doc, "AI", "일기와 음성을 범용 AI 모델 학습에 사용하지 않습니다.")

    add_heading(doc, "1. 서비스 제공을 위한 개인정보 처리 안내", 1)
    add_info_table(
        doc,
        [
            ("처리 목적", "회원 식별, 일기 작성·저장·조회·삭제, 기억 카드 제공, 서비스 이용 기록 관리"),
            ("처리 항목", "계정 식별정보, 닉네임, 일기 내용, 작성 일시, 인물·장소·감정 등 기억 태그, 서비스 이용 기록"),
            ("보유 기간", "회원 탈퇴 또는 이용자가 해당 기록을 삭제할 때까지. 법령에 따라 보존해야 하는 정보는 해당 기간 동안 분리 보관"),
            ("확인 거부 시", "서비스 제공에 필요한 최소 정보의 처리를 거부하면 회원 기능과 일기 저장 기능 이용이 어려울 수 있습니다."),
        ],
    )
    add_check_line(doc, "위 내용을 읽고 확인했습니다. (필수)", "")

    # Page 2: optional voice and sensitive information.
    add_heading(doc, "2. 음성 녹음 및 글 변환 동의", 1, page_break_before=True)
    add_body(doc, "선택 기능입니다. 동의하지 않아도 직접 글을 입력하여 일기를 작성할 수 있습니다.")
    add_info_table(
        doc,
        [
            ("이용 목적", "음성으로 일기 작성, 음성을 글로 변환, 이용자가 말한 내용 확인 및 수정"),
            ("처리 항목", "음성 녹음 파일, 녹음 일시·길이, 음성에서 변환된 글"),
            ("보유 기간", "원본 음성은 글 변환 후 삭제하는 것을 원칙으로 합니다. 원음 저장을 선택한 경우에는 개별 삭제 또는 회원 탈퇴 시까지, 변환된 글은 일기 삭제 또는 회원 탈퇴 시까지 보관합니다."),
            ("거부 시 영향", "음성 기능만 제한됩니다. 직접 글을 입력하는 일기 기능은 계속 이용할 수 있습니다."),
        ],
    )
    add_check_line(doc)

    add_heading(doc, "3. 건강·인지 관련 민감정보 처리 동의", 1)
    add_body(
        doc,
        "건강, 질병, 복약, 심리 또는 기억 상태를 일기에 남기면 민감한 정보가 포함될 수 있습니다. Haru는 동의한 경우에만 해당 정보를 개인화 기능에 활용합니다.",
    )
    add_info_table(
        doc,
        [
            ("이용 목적", "개인 맞춤형 기억 질문, 일상 루틴 지원, 반복 활동에 따른 Haru 지원 신호 제공"),
            ("처리 항목", "이용자가 직접 입력하거나 말한 건강·복약·심리·인지 관련 내용, 일상 활동 기록에서 생성된 지원 신호"),
            ("보유 기간", "해당 기록의 개별 삭제, 동의 철회 또는 회원 탈퇴 시까지"),
            ("거부 시 영향", "일반 일기는 이용할 수 있으나 건강·인지 개인화 기능과 지원 신호는 제공되지 않습니다."),
        ],
    )
    add_body(
        doc,
        "중요 안내: Haru의 지원 신호는 의료 진단이나 치료 결과가 아닙니다.",
        bold_prefix="중요 안내:",
        after=2,
    )
    add_check_line(doc, "민감정보 처리에 동의합니다", "동의하지 않습니다")

    # Page 3: family sharing and AI statement.
    add_heading(doc, "4. 지정 가족·보호자에게 정보 공유 동의", 1, page_break_before=True)
    add_body(
        doc,
        "가족 공유는 이용자가 특정 가족 또는 보호자를 직접 연결하고 공유 범위를 선택한 경우에만 시작됩니다.",
    )
    add_info_table(
        doc,
        [
            ("제공받는 사람", "이용자가 직접 지정하여 연결한 가족 또는 보호자 계정"),
            ("제공 목적", "함께 보기로 선택한 일기 확인, 추억 대화 준비, 일상 지원"),
            ("제공 항목", "이용자가 공유로 설정한 일기 제목·내용·기억 태그·작성 일시"),
            ("보유·이용 기간", "이용자가 공유를 중단하거나 가족 연결을 해제할 때까지"),
            ("거부 시 영향", "가족 공유 기능만 제한되며 개인 일기 기능은 계속 이용할 수 있습니다."),
        ],
    )
    add_check_line(doc, "지정 가족·보호자에게 정보 제공에 동의합니다", "동의하지 않습니다")

    add_heading(doc, "민감정보 공유 추가 동의", 2)
    add_body(
        doc,
        "건강·복약·심리·인지 관련 내용과 Haru 지원 신호는 아래 동의가 있는 경우에만 지정된 가족 또는 보호자에게 공유합니다. 원본 음성은 기본 공유 대상에 포함하지 않습니다.",
    )
    add_check_line(doc, "민감정보 제공에 동의합니다", "동의하지 않습니다")

    add_heading(doc, "5. AI 이용 안내", 1)
    add_callout(
        doc,
        "AI 학습 안 함",
        "오가사 LABs는 이용자의 일기, 음성 및 변환된 글을 범용 AI 모델 학습에 사용하지 않습니다.",
        fill=PALE_BLUE,
        accent=BLUE,
    )
    add_body(doc, "동의한 정보는 Haru 안에서 일기 정리, 기억 질문 생성, 개인화 기능 제공 목적으로만 처리합니다.")
    add_body(doc, "이용자의 정보를 광고 목적으로 판매하거나 제공하지 않습니다.")
    add_body(doc, "향후 AI 이용 목적이 달라지는 경우 변경 내용을 먼저 알리고 필요한 동의를 다시 받습니다.")

    # Page 4: user rights and third-person information notice.
    add_heading(doc, "6. 이용자의 권리", 1, page_break_before=True)
    rights = [
        ("열람", "저장된 개인정보와 일기를 확인할 수 있습니다."),
        ("정정", "잘못된 내용을 수정할 수 있습니다."),
        ("삭제", "일기·음성·변환된 글을 개별 삭제할 수 있습니다."),
        ("철회", "선택 동의를 철회하고 가족 공유를 중단할 수 있습니다."),
        ("이동·탈퇴", "개인정보 내보내기와 회원 탈퇴를 요청할 수 있습니다."),
    ]
    for label, detail in rights:
        add_easy_promise(doc, label, detail)
    add_body(doc, "이용 경로: Haru 설정 → 개인정보 및 동의 관리", bold_prefix="이용 경로:")
    add_body(doc, "개인정보 문의: 오가사 LABs (연락처는 운영 전 확정)", bold_prefix="개인정보 문의:")

    add_heading(doc, "7. 다른 사람의 개인정보에 관한 안내", 1)
    add_callout(
        doc,
        "기록 전 확인",
        "일기에는 다른 사람의 실명, 연락처, 주소, 병력 등 식별 가능한 정보를 적지 말고 별칭을 사용해 주세요. 다른 사람과 나눈 비공개 대화를 동의 없이 녹음하거나 공유하지 마세요.",
        fill=PALE_GOLD,
        accent=GOLD,
    )

    # Page 5: consent summary and signed confirmation with supplied stamp.
    add_heading(doc, "동의 확인", 1, page_break_before=True)
    add_body(doc, "본인은 위 내용을 읽고 이해했으며, 각 선택 항목에 대해 직접 동의 여부를 결정했습니다.")

    summary = doc.add_table(rows=4, cols=2)
    set_table_geometry(summary, [7000, 2360])
    set_table_borders(summary)
    choices = [
        ("서비스 제공을 위한 개인정보 처리 안내", "□ 확인"),
        ("음성 녹음 및 글 변환", "□ 동의  □ 미동의"),
        ("건강·인지 관련 민감정보 처리", "□ 동의  □ 미동의"),
        ("지정 가족·보호자 정보 제공", "□ 동의  □ 미동의"),
    ]
    for idx, (label, choice) in enumerate(choices):
        c1, c2 = summary.rows[idx].cells
        c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        c2.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if idx % 2 == 0:
            set_cell_shading(c1, "FAFBFC")
            set_cell_shading(c2, "FAFBFC")
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(label)
        set_run_font(r1, size=10.5, color=INK)
        p2 = c2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(choice)
        set_run_font(r2, size=10.5, bold=True, color=NAVY)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    signature = doc.add_table(rows=1, cols=2)
    set_table_geometry(signature, [4680, 4680])
    remove_table_borders(signature)
    left_cell, right_cell = signature.rows[0].cells
    for cell in (left_cell, right_cell):
        set_cell_margins(cell, top=120, bottom=80, start=120, end=120)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    lp = left_cell.paragraphs[0]
    lr = lp.add_run("이용자 확인")
    set_run_font(lr, size=11, bold=True, color=BLUE)
    for text in (
        "동의 일자: ______년 ____월 ____일",
        "이름 또는 닉네임: ____________________",
        "서명 또는 전자 확인: __________________",
    ):
        p = left_cell.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        r = p.add_run(text)
        set_run_font(r, size=10.5, color=INK)

    rp = right_cell.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = rp.add_run("개인정보처리자 확인")
    set_run_font(rr, size=11, bold=True, color=BLUE)
    op = right_cell.add_paragraph()
    op.alignment = WD_ALIGN_PARAGRAPH.CENTER
    op.paragraph_format.space_after = Pt(0)
    or1 = op.add_run("오가사 LABs\n대표 김현준  (인)")
    set_run_font(or1, size=11, bold=True, color=NAVY)
    stamp_p = right_cell.add_paragraph()
    stamp_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    stamp_p.paragraph_format.space_before = Pt(0)
    stamp_p.paragraph_format.space_after = Pt(0)
    stamp_run = stamp_p.add_run()
    stamp_run.add_picture(str(STAMP), width=Inches(0.92))
    drawings = stamp_run._element.xpath(".//wp:docPr")
    if drawings:
        drawings[0].set("descr", "오가사 LABs 대표 김현준 직인")
        drawings[0].set("title", "대표 직인")

    bottom = doc.add_paragraph()
    bottom.paragraph_format.space_before = Pt(8)
    bottom.paragraph_format.space_after = Pt(0)
    bottom.alignment = WD_ALIGN_PARAGRAPH.CENTER
    br = bottom.add_run("문서 버전  Demo Draft 0.1    |    시행 예정일  별도 공지")
    set_run_font(br, size=8.5, color=MUTED)

    return doc


if __name__ == "__main__":
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT)
    print(OUTPUT)
