from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"C:\project\saerok-memory")
OUTPUT = ROOT / "docs" / "Haru_개인정보_처리_동의서_데모.docx"
STAMP = Path(r"C:\Users\mnb92\Downloads\도장.png")

# compact_reference_guide base with named override: official_korean_two_page.
BODY_FONT = "HCR Batang"
BOLD_FONT = "HCR Batang Bold"
TITLE_FONT = "HY궁서B"
BLACK = "000000"
GRAY_FILL = "F2F2F2"
HEADER_FILL = "E7E7E7"
NOTE_FILL = "F7F7F7"
BORDER = "777777"

# A4 with 1.4 cm left/right margins: 18.2 cm / 10,318 DXA usable width.
TABLE_WIDTH_DXA = 10318
TABLE_INDENT_DXA = 70


def set_run_font(
    run,
    *,
    font: str = BODY_FONT,
    size: float = 9.2,
    bold: bool = False,
    italic: bool = False,
) -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(BLACK)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{key}"), font)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top=30, start=70, bottom=30, end=70) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, *, color: str = BORDER, size: str = "4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def set_table_geometry(table, widths_dxa: list[int], *, indent_dxa: int = TABLE_INDENT_DXA) -> None:
    assert sum(widths_dxa) == TABLE_WIDTH_DXA
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run("쪽 ")
    set_run_font(run, size=8.0)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_para(
    container,
    text: str,
    *,
    size: float = 9.2,
    font: str = BODY_FONT,
    bold: bool = False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    before: float = 0,
    after: float = 2,
    line: float = 1.03,
    keep_with_next: bool = False,
    page_break_before: bool = False,
):
    paragraph = container.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.keep_with_next = keep_with_next
    paragraph.paragraph_format.page_break_before = page_break_before
    run = paragraph.add_run(text)
    set_run_font(run, font=font, size=size, bold=bold)
    return paragraph


def add_section_heading(doc: Document, text: str, *, page_break_before: bool = False) -> None:
    add_para(
        doc,
        text,
        size=11.3,
        font=BOLD_FONT,
        bold=True,
        before=4,
        after=3,
        line=1.0,
        keep_with_next=True,
        page_break_before=page_break_before,
    )


def format_cell_text(cell, text: str, *, bold: bool = False, size: float = 8.4, center: bool = False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    set_run_font(run, font=BOLD_FONT if bold else BODY_FONT, size=size, bold=bold)


def add_note_box(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [TABLE_WIDTH_DXA], indent_dxa=90)
    set_table_borders(table, color="AAAAAA", size="3")
    cell = table.cell(0, 0)
    set_cell_shading(cell, NOTE_FILL)
    set_cell_margins(cell, top=45, bottom=45, start=90, end=90)
    format_cell_text(cell, text, size=8.5)
    add_para(doc, "", size=1, after=1)


def add_primary_consent_table(doc: Document) -> None:
    widths = [1300, 1880, 2800, 3020, 1318]
    table = doc.add_table(rows=4, cols=5)
    set_table_geometry(table, widths)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    headers = ("구분", "처리 목적", "처리 항목", "보유 기간 및 거부 영향", "확인")
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, HEADER_FILL)
        format_cell_text(cell, text, bold=True, size=8.5, center=True)

    rows = [
        (
            "서비스 제공\n(필수 확인)",
            "회원 식별, 일기 저장·조회·삭제, 기억 카드 및 이용 기록 관리",
            "계정 식별정보, 닉네임, 일기 내용·작성일시·기억 태그, 이용 기록",
            "개별 삭제 또는 회원 탈퇴 시까지. 거부하면 회원·일기 저장 기능 이용이 어려울 수 있음",
            "□ 확인",
        ),
        (
            "음성 녹음·전사\n(선택)",
            "음성 일기 작성 및 음성을 글로 변환",
            "원본 음성, 녹음 일시·길이, 변환된 글",
            "원음은 전사 후 삭제가 원칙. 저장 선택 시 개별 삭제·탈퇴까지. 거부해도 글 일기 이용 가능",
            "□ 동의\n□ 미동의",
        ),
        (
            "건강·인지 정보\n(선택)",
            "맞춤 기억 질문, 일상 루틴 및 Haru 지원 신호 제공",
            "건강·복약·심리·인지 관련 내용 및 활동 기록에서 생성된 지원 신호",
            "개별 삭제·동의 철회·탈퇴 시까지. 거부해도 일반 일기 이용 가능",
            "□ 동의\n□ 미동의",
        ),
    ]
    for row_index, row_data in enumerate(rows, start=1):
        for col_index, text in enumerate(row_data):
            cell = table.rows[row_index].cells[col_index]
            if col_index == 0:
                set_cell_shading(cell, GRAY_FILL)
            format_cell_text(
                cell,
                text,
                bold=col_index in (0, 4),
                size=8.25 if col_index not in (0, 4) else 8.4,
                center=col_index in (0, 4),
            )
    add_para(doc, "※ Haru 지원 신호는 의료 진단이나 치료 결과가 아닙니다.", size=8.3, after=3)


def add_family_table(doc: Document) -> None:
    widths = [1450, 2750, 3910, 2208]
    table = doc.add_table(rows=3, cols=4)
    set_table_geometry(table, widths)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    headers = ("구분", "제공받는 사람·목적", "제공 항목·기간", "동의")
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, HEADER_FILL)
        format_cell_text(cell, text, bold=True, size=8.5, center=True)

    rows = [
        (
            "일반정보 공유",
            "이용자가 지정한 가족·보호자 / 공유 일기 확인과 추억 대화 지원",
            "공유 설정한 일기 내용·기억 태그·작성일시 / 연결 해제 또는 공유 중단까지",
            "□ 동의  □ 미동의",
        ),
        (
            "민감정보 공유",
            "이용자가 지정한 가족·보호자 / 일상 지원",
            "건강·복약·심리·인지 내용과 지원 신호 / 연결 해제 또는 공유 중단까지. 원본 음성 제외",
            "□ 동의  □ 미동의",
        ),
    ]
    for row_index, row_data in enumerate(rows, start=1):
        for col_index, text in enumerate(row_data):
            cell = table.rows[row_index].cells[col_index]
            if col_index == 0:
                set_cell_shading(cell, GRAY_FILL)
            format_cell_text(
                cell,
                text,
                bold=col_index in (0, 3),
                size=8.2 if col_index not in (0, 3) else 8.35,
                center=col_index in (0, 3),
            )
    add_para(doc, "※ 가족 공유는 이용자가 수신자와 범위를 직접 선택한 경우에만 시작됩니다.", size=8.3, after=2)


def add_summary_table(doc: Document) -> None:
    widths = [6900, 3418]
    table = doc.add_table(rows=5, cols=2)
    set_table_geometry(table, widths)
    set_table_borders(table)
    rows = [
        ("서비스 제공을 위한 개인정보 처리 안내", "□ 확인"),
        ("음성 녹음 및 글 변환", "□ 동의  □ 미동의"),
        ("건강·인지 관련 민감정보 처리", "□ 동의  □ 미동의"),
        ("지정 가족·보호자 일반정보 제공", "□ 동의  □ 미동의"),
        ("지정 가족·보호자 민감정보 제공", "□ 동의  □ 미동의"),
    ]
    for row_index, (label, choice) in enumerate(rows):
        left, right = table.rows[row_index].cells
        if row_index % 2 == 0:
            set_cell_shading(left, GRAY_FILL)
            set_cell_shading(right, GRAY_FILL)
        format_cell_text(left, label, size=8.7)
        format_cell_text(right, choice, bold=True, size=8.7, center=True)


def style_document(doc: Document) -> None:
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.05)
    section.bottom_margin = Cm(1.05)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)
    section.header_distance = Cm(0.45)
    section.footer_distance = Cm(0.45)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.03

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("Haru 개인정보 처리 안내 및 동의서 (계속)")
    set_run_font(hr, font=BOLD_FONT, size=8.2, bold=True)
    hp.add_run("\t")
    dr = hp.add_run("데모용 초안")
    set_run_font(dr, font=BODY_FONT, size=8.2)
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(18.2), WD_TAB_ALIGNMENT.RIGHT)

    for footer in (section.footer, section.first_page_footer):
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fp.paragraph_format.space_before = Pt(0)
        fp.paragraph_format.space_after = Pt(0)
        fr = fp.add_run("오가사 LABs  |  대표 김현준")
        set_run_font(fr, size=8.0)
        fp.add_run("\t")
        fp.paragraph_format.tab_stops.add_tab_stop(Cm(18.2), WD_TAB_ALIGNMENT.RIGHT)
        add_page_field(fp)


def build_document() -> Document:
    if not STAMP.exists():
        raise FileNotFoundError(STAMP)

    doc = Document()
    style_document(doc)
    props = doc.core_properties
    props.title = "Haru 개인정보 처리 안내 및 동의서 - 데모"
    props.subject = "개인 일기 및 음성 기록 서비스 개인정보 동의서 데모"
    props.author = "오가사 LABs"
    props.keywords = "Haru, 개인정보, 동의서, 데모"

    add_para(
        doc,
        "개인정보 처리 안내 및 동의서",
        size=19.0,
        font=TITLE_FONT,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=0,
        after=2,
        line=1.0,
        keep_with_next=True,
    )
    add_para(
        doc,
        "개인 일기·음성 기록 서비스  |  데모용 초안",
        size=9.0,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=1,
        line=1.0,
        keep_with_next=True,
    )
    add_para(
        doc,
        "개인정보처리자  오가사 LABs    대표자  김현준    소재 국가  대한민국",
        size=8.8,
        font=BOLD_FONT,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=3,
        line=1.0,
    )
    add_note_box(doc, "본 문서는 시연용 초안입니다. 실제 운영 전 보유기간, 수탁업체 및 문의처를 확정해야 합니다.")

    add_section_heading(doc, "1. 개인정보 수집·이용 및 선택 동의")
    add_primary_consent_table(doc)

    add_section_heading(doc, "2. 지정 가족·보호자 정보 제공 동의")
    add_family_table(doc)

    add_section_heading(doc, "3. AI 이용 및 이용자 권리", page_break_before=True)
    add_note_box(
        doc,
        "AI 이용 안내  |  오가사 LABs는 일기·음성·변환된 글을 범용 AI 모델 학습에 사용하지 않습니다. 동의한 정보는 일기 정리, 기억 질문 및 개인화 기능에만 활용하며 광고 목적으로 판매·제공하지 않습니다.",
    )
    add_para(
        doc,
        "이용자 권리  |  이용자는 개인정보 열람·정정·삭제, 선택 동의 철회, 가족 공유 중단, 내보내기 및 회원 탈퇴를 요청할 수 있습니다. 이용 경로는 Haru 설정의 개인정보 및 동의 관리이며, 문의처는 오가사 LABs가 운영 전 확정합니다.",
        size=8.8,
        after=3,
        line=1.05,
    )
    add_para(
        doc,
        "다른 사람의 정보  |  일기에는 타인의 실명·연락처·주소·병력 등 식별정보 대신 별칭을 사용하고, 비공개 대화를 동의 없이 녹음하거나 공유하지 마세요.",
        size=8.8,
        after=4,
        line=1.05,
    )

    add_section_heading(doc, "4. 최종 동의 확인")
    add_para(
        doc,
        "본인은 위 내용을 읽고 이해했으며 각 선택 항목의 동의 여부를 직접 결정했습니다.",
        size=8.8,
        after=2,
    )
    add_summary_table(doc)

    add_para(doc, "", size=1, after=1)
    signature = doc.add_table(rows=1, cols=2)
    set_table_geometry(signature, [5159, 5159])
    remove_table_borders(signature)
    left, right = signature.rows[0].cells
    for cell in (left, right):
        set_cell_margins(cell, top=55, bottom=20, start=70, end=70)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    format_cell_text(left, "이용자 확인", bold=True, size=9.4)
    left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for text in (
        "동의 일자: ______년 ____월 ____일",
        "이름 또는 닉네임: ____________________",
        "서명 또는 전자 확인: __________________",
    ):
        p = left.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        set_run_font(r, size=8.9)

    format_cell_text(right, "개인정보처리자 확인", bold=True, size=9.4)
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    operator = right.add_paragraph()
    operator.alignment = WD_ALIGN_PARAGRAPH.CENTER
    operator.paragraph_format.space_before = Pt(0)
    operator.paragraph_format.space_after = Pt(0)
    operator_run = operator.add_run("오가사 LABs    대표 김현준  (인)")
    set_run_font(operator_run, font=BOLD_FONT, size=9.4, bold=True)
    stamp_p = right.add_paragraph()
    stamp_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    stamp_p.paragraph_format.space_before = Pt(0)
    stamp_p.paragraph_format.space_after = Pt(0)
    stamp_run = stamp_p.add_run()
    stamp_run.add_picture(str(STAMP), width=Inches(0.72))
    doc_props = stamp_run._element.xpath(".//wp:docPr")
    if doc_props:
        doc_props[0].set("descr", "오가사 LABs 대표 김현준 직인")
        doc_props[0].set("title", "대표 직인")

    add_para(
        doc,
        "문서 버전  Demo Draft 0.2  |  시행 예정일  별도 공지",
        size=7.8,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=1,
        after=0,
        line=1.0,
    )
    return doc


if __name__ == "__main__":
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT)
    print(OUTPUT)
